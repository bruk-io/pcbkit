"""Integration: pcbkit.fab against a real KiCad, on the tiny board.

``export_fab`` runs for real on a throwaway project built around tests/tiny_board.py
(two 0603 resistors on a 30 x 20 mm board): kicad-cli plots the Gerbers and drill files,
writes the positions and the PDFs, rsvg-convert draws the assembly drawing, and pcbnew
sets the origin and finds the body centres. The files are then read back and `pcbkit
quote` is run on them. A planted mistake (the origin not set) shows the centroid check
can fail.

These tests need KiCad's own Python, where ``import pcbnew`` works, and kicad-cli: see
tests/integration/test_kicad_core.py or .claude/CLAUDE.md for how to make .venv-kicad.
In any other Python they are skipped, for the reason given just below.
"""

from __future__ import annotations

import csv
import json
import re
import shutil
import zipfile
from collections.abc import Iterator
from pathlib import Path

import click
import pytest
from openpyxl import load_workbook

pcbnew = pytest.importorskip(
    "pcbnew",
    reason="pcbnew only imports under KiCad's own Python: see tests/integration/"
    "test_kicad_core.py or .claude/CLAUDE.md for how to make .venv-kicad",
)

from pcbkit.fab import FabResult, bom, export_fab, gerbers, pcbway  # noqa: E402
from pcbkit.kicad.sexp import find, parse  # noqa: E402
from pcbkit.project import Project, load_project  # noqa: E402
from tests import tiny_board  # noqa: E402
from tests.board_files import restored_imports, write_file  # noqa: E402

pytestmark = pytest.mark.kicad

PNG = b"\x89PNG\r\n\x1a\n"

TOML = """\
[board]
stem = "tiny"
title = "Tiny"
rev = "A"
fab_name = "Tiny_revA"
"""

DESIGN = """\
from pcbkit.design import R

R("R1", "10k", "NET_A", "NET_B", "Block")
R("R2", "10k", "NET_B", "GND", "Block")
"""

LAYOUT = "W = 30.0\nH = 20.0\n"

# The nine plotted layers, the job file, and the plated drill file and its map. (Which
# unplated files exist depends on whether the board has unplated holes.)
GERBER_NAMES = {
    "tiny-F_Cu.gbr",
    "tiny-B_Cu.gbr",
    "tiny-F_Paste.gbr",
    "tiny-B_Paste.gbr",
    "tiny-F_Silkscreen.gbr",
    "tiny-B_Silkscreen.gbr",
    "tiny-F_Mask.gbr",
    "tiny-B_Mask.gbr",
    "tiny-Edge_Cuts.gbr",
    "tiny-job.gbrjob",
    "tiny-PTH.drl",
    "tiny-PTH-drl_map.gbr",
}


@pytest.fixture(autouse=True)
def clean_imports() -> Iterator[None]:
    """Undo what loading a project's modules does to sys.path and sys.modules."""
    with restored_imports():
        yield


@pytest.fixture(scope="module")
def tiny(tmp_path_factory: pytest.TempPathFactory) -> tiny_board.TinyBoard:
    """Build the tiny board once; projects copy it."""
    return tiny_board.build(tmp_path_factory.mktemp("tiny"))


def make_project(
    root: Path,
    tiny: tiny_board.TinyBoard,
    design: str = DESIGN,
    layout: str = LAYOUT,
) -> Project:
    """Make a project in ``root`` around a copy of the tiny board and return it."""
    kicad = root / "kicad"
    kicad.mkdir(parents=True)
    for path in (tiny.pcb, tiny.sch):
        shutil.copy(path, kicad / path.name)
    write_file(root / "pcbkit.toml", TOML)
    write_file(root / "design.py", design)
    write_file(root / "layout.py", layout)
    return load_project(root)


class Exported:
    """A project and the result of exporting it, for tests that only read the files."""

    def __init__(self, project: Project, result: FabResult) -> None:
        self.project = project
        self.result = result
        self.fab = project.out_dir / "fab"
        self.docs = project.out_dir / "docs"


@pytest.fixture(scope="module")
def exported(
    tmp_path_factory: pytest.TempPathFactory, tiny: tiny_board.TinyBoard
) -> Exported:
    """Export the tiny project once, without renders, and share the files."""
    with restored_imports():
        project = make_project(tmp_path_factory.mktemp("project"), tiny)
        return Exported(project, export_fab(project, render=False))


# --- what is written -----------------------------------------------------------------


def test_the_export_writes_the_fab_files_and_the_documents(exported: Exported) -> None:
    assert {p.name for p in (exported.fab / "gerbers").iterdir()} >= GERBER_NAMES
    for name in (
        "Tiny_revA_gerbers.zip",
        "Tiny_revA_BOM.csv",
        "Tiny_revA_BOM.xlsx",
        "Tiny_revA_centroid.csv",
    ):
        assert (exported.fab / name).is_file(), name
    assert {p.name for p in exported.docs.iterdir()} == {
        "Tiny_revA_schematic.pdf",
        "Tiny_revA_assembly_top.pdf",
        "Tiny_revA_assembly_top.png",
        "Tiny_revA_top_copper.pdf",
        "Tiny_revA_bottom_copper.pdf",
    }


def test_the_result_lists_every_file_and_the_bom_size(exported: Exported) -> None:
    result = exported.result
    assert (result.bom_lines, result.total_parts) == (1, 2)
    assert all(path.is_file() for path in result.files)
    on_disk = {p for p in exported.project.out_dir.rglob("*") if p.is_file()}
    assert set(result.files) == on_disk
    assert result.without_mpn == []


def test_the_zip_holds_exactly_the_files_in_the_gerber_folder(
    exported: Exported,
) -> None:
    folder = exported.fab / "gerbers"
    with zipfile.ZipFile(exported.fab / "Tiny_revA_gerbers.zip") as archive:
        assert sorted(archive.namelist()) == sorted(p.name for p in folder.iterdir())
        for name in archive.namelist():
            assert archive.read(name) == (folder / name).read_bytes()


def test_the_documents_are_real_pdfs_and_a_png(exported: Exported) -> None:
    for name in ("schematic", "top_copper", "bottom_copper", "assembly_top"):
        data = (exported.docs / f"Tiny_revA_{name}.pdf").read_bytes()
        assert data.startswith(b"%PDF-"), name
    png = (exported.docs / "Tiny_revA_assembly_top.png").read_bytes()
    assert png.startswith(PNG)


def test_the_bom_is_made_from_the_design_and_the_default_resistor_table(
    exported: Exported,
) -> None:
    with open(exported.fab / "Tiny_revA_BOM.csv", newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert rows[0] == list(bom.HEADER)
    assert rows[1] == [
        "1",
        "2",
        "R1,R2",
        "Yageo",
        "RC0603FR-0710KL",
        "10k",
        "Resistor 10k 0603 1%",
        "R_0603_1608Metric",
        "SMD",
    ]
    assert len(rows) == 2


def test_the_workbook_has_the_same_rows(exported: Exported) -> None:
    sheet = load_workbook(exported.fab / "Tiny_revA_BOM.xlsx")["BOM"]
    cells = [list(row) for row in sheet.iter_rows(values_only=True)]
    assert cells[0] == list(bom.HEADER)
    assert cells[1][:5] == [1, 2, "R1,R2", "Yageo", "RC0603FR-0710KL"]


# --- the centroid and the origin -----------------------------------------------------


def centroid_of(folder: Path) -> dict[str, dict[str, str]]:
    """Return the centroid file's rows by designator."""
    with open(folder / "Tiny_revA_centroid.csv", newline="", encoding="utf-8") as f:
        return {row["Designator"]: row for row in csv.DictReader(f)}


def test_the_centroid_is_measured_from_the_bottom_left_corner_with_y_up(
    exported: Exported,
) -> None:
    rows = centroid_of(exported.fab)
    assert set(rows) == {"R1", "R2"}
    # tiny_board places R1 at (8, 10) and R2 at (22, 10) from the top-left of a board
    # 20 mm tall, so from the bottom-left they are 10 mm up.
    for ref, x in (("R1", 8.0), ("R2", 22.0)):
        assert float(rows[ref]["Mid X (mm)"]) == pytest.approx(x, abs=1e-6)
        assert float(rows[ref]["Mid Y (mm)"]) == pytest.approx(10.0, abs=1e-6)
        assert rows[ref]["Layer"] == "Top"
        assert rows[ref]["Footprint"] == "R_0603_1608Metric"


def test_the_export_puts_the_aux_origin_at_the_bottom_left_corner(
    exported: Exported,
) -> None:
    pcb = exported.project.kicad_dir / "tiny.kicad_pcb"
    setup = find(parse(pcb.read_text(encoding="utf-8")), "setup")
    assert setup is not None
    origin = find(setup, "aux_axis_origin")
    assert origin is not None
    assert (float(origin[1]), float(origin[2])) == (50.0, 70.0)  # (50, 50 + H)


def test_an_export_that_forgot_the_origin_would_fail_the_centroid_check(
    tmp_path: Path, tiny: tiny_board.TinyBoard, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Plant the mistake: skip set_origin, and the positions come out wrong."""
    project = make_project(tmp_path, tiny)
    monkeypatch.setattr(gerbers, "set_origin", lambda pcb, height: None)
    export_fab(project, render=False)
    x = float(centroid_of(project.out_dir / "fab")["R1"]["Mid X (mm)"])
    assert x != pytest.approx(8.0, abs=0.01)


def test_a_wrong_origin_already_in_the_board_is_replaced(
    tmp_path: Path, tiny: tiny_board.TinyBoard
) -> None:
    project = make_project(tmp_path, tiny)
    pcb = project.kicad_dir / "tiny.kicad_pcb"
    board = pcbnew.LoadBoard(str(pcb))
    board.GetDesignSettings().SetAuxOrigin(pcbnew.VECTOR2I(pcbnew.FromMM(61), 7))
    pcbnew.SaveBoard(str(pcb), board)
    export_fab(project, render=False)
    assert float(centroid_of(project.out_dir / "fab")["R2"]["Mid X (mm)"]) == (
        pytest.approx(22.0, abs=1e-6)
    )


def test_a_part_whose_anchor_is_off_its_body_gets_the_body_centre(
    tmp_path: Path, tiny: tiny_board.TinyBoard
) -> None:
    project = make_project(tmp_path, tiny)
    pcb = project.kicad_dir / "tiny.kicad_pcb"
    board = pcbnew.LoadBoard(str(pcb))
    shape = pcbnew.PCB_SHAPE(board.FindFootprintByReference("R2"))
    shape.SetShape(pcbnew.SHAPE_T_SEGMENT)
    shape.SetLayer(pcbnew.F_Fab)
    shape.SetWidth(pcbnew.FromMM(0.1))
    footprint = board.FindFootprintByReference("R2")
    anchor = footprint.GetPosition()
    shape.SetStart(pcbnew.VECTOR2I(anchor.x + pcbnew.FromMM(3), anchor.y))
    shape.SetEnd(pcbnew.VECTOR2I(anchor.x + pcbnew.FromMM(4), anchor.y))
    footprint.Add(shape)
    pcbnew.SaveBoard(str(pcb), board)
    export_fab(project, render=False)
    rows = centroid_of(project.out_dir / "fab")
    # The stock outline runs -0.8 to 0.8 mm; the extra segment reaches 4 mm, so the body
    # centre sits 1.6 mm along from the anchor. The text is four decimals, not six.
    assert rows["R2"]["Mid X (mm)"] == "23.6000"
    assert rows["R2"]["Mid Y (mm)"] == "10.0000"
    assert rows["R1"]["Mid X (mm)"] == "8.000000"  # untouched: kicad-cli's own text


STACKUP = """\
\t\t(stackup
\t\t\t(layer "F.SilkS" (type "Top Silk Screen"))
\t\t\t(layer "F.Paste" (type "Top Solder Paste"))
\t\t\t(layer "F.Mask" (type "Top Solder Mask") (thickness 0.01))
\t\t\t(layer "F.Cu" (type "copper") (thickness {copper}))
\t\t\t(layer "dielectric 1" (type "core") (thickness 1.4) (material "FR4"))
\t\t\t(layer "B.Cu" (type "copper") (thickness {copper}))
\t\t\t(layer "B.Mask" (type "Bottom Solder Mask") (thickness 0.01))
\t\t\t(layer "B.Paste" (type "Bottom Solder Paste"))
\t\t\t(layer "B.SilkS" (type "Bottom Silk Screen"))
\t\t\t(copper_finish "HAL lead-free")
\t\t\t(dielectric_constraints no)
\t\t)
"""


def record_stackup(pcb: Path, copper: float) -> None:
    """Add a stackup with ``copper`` mm outer copper and a finish to a saved board."""
    text = pcb.read_text(encoding="utf-8")
    at = text.index("\n", text.index("(setup")) + 1
    pcb.write_text(
        text[:at] + STACKUP.format(copper=copper) + text[at:], encoding="utf-8"
    )


def test_a_stackup_recorded_in_the_board_survives_the_origin_save(
    tmp_path: Path, tiny: tiny_board.TinyBoard
) -> None:
    """Saving the board to set its origin must not put the default copper back."""
    project = make_project(tmp_path, tiny)
    pcb = project.kicad_dir / "tiny.kicad_pcb"
    record_stackup(pcb, 0.07)
    export_fab(project, render=False)
    with zipfile.ZipFile(project.out_dir / "fab" / "Tiny_revA_gerbers.zip") as archive:
        job = json.loads(archive.read("tiny-job.gbrjob"))
    copper = [m["Thickness"] for m in job["MaterialStackup"] if m["Type"] == "Copper"]
    assert copper == [0.07, 0.07]
    assert job["GeneralSpecs"]["Finish"] == "HAL lead-free"
    saved = pcb.read_text(encoding="utf-8")
    assert (
        len(
            re.findall(
                r'\(layer "[FB]\.Cu"\s*\(type "copper"\)\s*\(thickness 0\.07\)', saved
            )
        )
        == 2
    )


def test_the_quote_warns_when_the_gerber_copper_is_not_the_projects(
    tmp_path: Path, tiny: tiny_board.TinyBoard
) -> None:
    project = make_project(tmp_path, tiny)  # [stackup] copper_mm defaults to 0.035
    record_stackup(project.kicad_dir / "tiny.kicad_pcb", 0.07)
    export_fab(project, render=False)
    quote = pcbway.build_quote(project)
    assert quote.board.copper_mm == (0.07, 0.07)
    assert quote.board.finish == "HAL lead-free"
    [warning] = quote.warnings
    assert warning.startswith("the Gerbers say 0.07, 0.07 mm copper but [stackup]")
    assert "[stackup] copper_mm is 0.035" in warning


# --- a second export, and what stops one --------------------------------------------


def test_a_second_export_leaves_no_file_of_the_first_and_none_outside_its_folders(
    tmp_path: Path, tiny: tiny_board.TinyBoard
) -> None:
    project = make_project(tmp_path, tiny)
    export_fab(project, render=False)
    (project.out_dir / "fab" / "gerbers" / "stale-F_Cu.gbr").write_text("old")
    (project.out_dir / "docs" / "stale.pdf").write_text("old")
    (project.out_dir / "reports").mkdir()
    (project.out_dir / "reports" / "keep.txt").write_text("mine")
    export_fab(project, render=False)
    with zipfile.ZipFile(project.out_dir / "fab" / "Tiny_revA_gerbers.zip") as archive:
        assert "stale-F_Cu.gbr" not in archive.namelist()
    assert not (project.out_dir / "fab" / "gerbers" / "stale-F_Cu.gbr").exists()
    assert not (project.out_dir / "docs" / "stale.pdf").exists()
    assert (project.out_dir / "reports" / "keep.txt").read_text() == "mine"


def test_a_missing_board_says_what_to_run(
    tmp_path: Path, tiny: tiny_board.TinyBoard
) -> None:
    project = make_project(tmp_path, tiny)
    (project.kicad_dir / "tiny.kicad_pcb").unlink()
    with pytest.raises(click.ClickException, match=r"tiny\.kicad_pcb not found"):
        export_fab(project, render=False)
    assert not (project.out_dir / "fab").exists()  # nothing was emptied or written


@pytest.mark.parametrize("layout", ["W = 30.0\n", "W = 30.0\nH = 'tall'\n", "H = 0\n"])
def test_a_layout_without_a_usable_height_is_an_error(
    tmp_path: Path, tiny: tiny_board.TinyBoard, layout: str
) -> None:
    project = make_project(tmp_path, tiny, layout=layout)
    with pytest.raises(click.ClickException, match="layout.py should define H"):
        export_fab(project, render=False)


def test_a_part_in_the_design_but_not_on_the_board_is_an_error(
    tmp_path: Path, tiny: tiny_board.TinyBoard
) -> None:
    design = DESIGN + 'R("R3", "10k", "NET_A", "GND", "Block")\n'
    project = make_project(tmp_path, tiny, design=design)
    with pytest.raises(bom.BomError, match="R3 are in design.py but not on the board"):
        export_fab(project, render=False)


def test_a_part_that_cannot_be_ordered_is_reported_not_fatal(
    tmp_path: Path, tiny: tiny_board.TinyBoard
) -> None:
    design = DESIGN.replace('"10k", "NET_B", "GND"', '"4.3k", "NET_B", "GND"')
    project = make_project(tmp_path, tiny, design=design)
    result = export_fab(project, render=False)
    [problem] = result.without_mpn
    assert (problem.ref, problem.mpn) == ("R2", "0603 4.3k 1%")
    assert result.bom_lines == 2  # the BOM is still written


def test_a_projects_bom_py_is_used(tmp_path: Path, tiny: tiny_board.TinyBoard) -> None:
    project = make_project(tmp_path, tiny)
    write_file(
        project.root / "bom.py",
        """
        REF_OVERRIDE = {"R1": {"mpn": "MY-PART", "mfr": "Me", "qty": 3}}
        NOT_IN_BOM = ["R9 is imaginary."]
        """,
    )
    result = export_fab(project, render=False)
    assert (result.bom_lines, result.total_parts) == (2, 4)
    sheet = load_workbook(project.out_dir / "fab" / "Tiny_revA_BOM.xlsx")["BOM"]
    cells = [list(row)[:5] for row in sheet.iter_rows(values_only=True)]
    assert cells[1] == [1, 3, "R1", "Me", "MY-PART"]
    assert cells[-1] == ["R9 is imaginary.", None, None, None, None]


# --- the renders and the quote -------------------------------------------------------


def test_the_renders_are_made_only_when_asked_for(
    tmp_path: Path, tiny: tiny_board.TinyBoard
) -> None:
    project = make_project(tmp_path, tiny)
    result = export_fab(project, render=True)
    for view in ("iso", "top", "bottom"):
        path = project.out_dir / "docs" / f"Tiny_revA_render_{view}.png"
        assert path.read_bytes().startswith(PNG), view
        assert path in result.files


def test_quote_reads_what_the_real_export_wrote(exported: Exported) -> None:
    quote = pcbway.build_quote(exported.project, assembled=2, fab_qty=5)
    board = quote.board
    assert board.layers == 2
    assert (board.width_mm, board.height_mm) == (30.0, 20.0)
    assert board.thickness_mm == pytest.approx(1.6)
    assert board.copper_mm == (0.035, 0.035)
    assert board.min_hole_mm == pytest.approx(0.4)  # the one via
    assert board.min_track_mm and board.min_track_mm > 0
    assert quote.warnings == ()
    assembly = quote.assembly
    assert assembly is not None
    assert (assembly.unique_parts, assembly.smd_placements) == (1, 2)
    assert assembly.through_hole_refs == ()
    text = pcbway.format_quote(quote)
    assert "Board size           30 x 20 mm" in text
