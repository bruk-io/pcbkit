"""Unit tests for pcbkit.fab.bom: grouping, sorting, overrides, completeness, files.

Every part here is made up. The tests build the parts the way a design.py does, and run
the real grouping, sorting and file writers on them; nothing is mocked.
"""

from __future__ import annotations

import csv
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from pcbkit import design
from pcbkit.design import Part
from pcbkit.fab import bom
from pcbkit.fab.bom import BomError, Group, Overrides
from tests.board_files import restored_imports, write_file

R0603 = "Resistor_SMD:R_0603_1608Metric"
C0603 = "Capacitor_SMD:C_0603_1608Metric"
HEADER3 = "Connector_PinHeader_2.54mm:PinHeader_1x03_P2.54mm_Vertical"


def make(
    ref: str,
    value: str = "10k",
    *,
    sym: str = "Device:R",
    fp: str = R0603,
    mfr: str = "",
    mpn: str | None = None,
    desc: str = "",
    dnp: bool = False,
    in_bom: bool = True,
) -> Part:
    """Return a part; a resistor given no ``mpn`` gets the R helper's stand-in."""
    if mpn is None:
        mpn = bom.stand_in_mpn(value) if sym == "Device:R" else ""
    return Part(
        ref=ref,
        sym=sym,
        value=value,
        fp=fp,
        pins={"1": "A", "2": "B"},
        mfr=mfr,
        mpn=mpn,
        desc=desc or f"{sym} {value}",
        block="Block",
        dnp=dnp,
        bom=in_bom,
    )


def lines_of(
    parts: list[Part], through_hole: set[str] | None = None, **kw: Any
) -> list[bom.BomLine]:
    """Return ``bom_lines`` for parts, with an Overrides built from ``kw``."""
    return bom.bom_lines(parts, through_hole or set(), Overrides(**kw))


@pytest.fixture(autouse=True)
def clean_imports() -> Iterator[None]:
    """Undo what loading a project's bom.py does to sys.path and sys.modules."""
    with restored_imports():
        yield


# --- grouping -----------------------------------------------------------------------


def test_parts_with_the_same_part_number_and_footprint_share_a_line() -> None:
    lines = lines_of([make("R2"), make("R1"), make("R3")])
    assert len(lines) == 1
    assert lines[0].refs == ("R1", "R2", "R3")
    assert lines[0].qty == 3
    assert lines[0].row()[:3] == [1, 3, "R1,R2,R3"]


def test_a_line_takes_its_text_from_the_first_part_declared_not_the_first_sorted() -> (
    None
):
    lines = lines_of(
        [
            make("R9", desc="declared first", mfr="Maker A"),
            make("R4", desc="declared second", mfr="Maker B"),
        ]
    )
    assert lines[0].refs == ("R4", "R9")
    assert lines[0].desc == "declared first"
    assert lines[0].mfr == "Yageo"  # the table's maker: both are stand-ins


def test_the_same_part_number_in_two_footprints_is_two_lines() -> None:
    parts = [
        make("C1", "100n", sym="Device:C", fp=C0603, mpn="CAP-1"),
        make(
            "C2",
            "100n",
            sym="Device:C",
            fp="Capacitor_SMD:C_0805_2012Metric",
            mpn="CAP-1",
        ),
    ]
    lines = lines_of(parts)
    assert [line.footprint for line in lines] == [
        "C_0603_1608Metric",
        "C_0805_2012Metric",
    ]


def test_parts_with_no_part_number_are_grouped_by_value_and_footprint() -> None:
    parts = [
        make("C1", "100n", sym="Device:C", fp=C0603),
        make("C2", "100n", sym="Device:C", fp=C0603),
        make("C3", "1u", sym="Device:C", fp=C0603),
    ]
    lines = lines_of(parts)
    assert [(line.value, line.refs) for line in lines] == [
        ("100n", ("C1", "C2")),
        ("1u", ("C3",)),
    ]
    assert all(line.mpn == "" for line in lines)


def test_flags_dnp_parts_and_parts_kept_out_of_the_bom_are_left_out() -> None:
    parts = [
        make("R1"),
        make("R2", dnp=True),
        make("R3", in_bom=False),
        make("#FLG01", sym="power:PWR_FLAG", fp="", mpn="X"),
    ]
    assert [line.refs for line in lines_of(parts)] == [("R1",)]


def test_a_footprint_without_a_library_name_is_taken_whole() -> None:
    assert bom.footprint_name(make("R1", fp="R_Plain")) == "R_Plain"
    assert bom.footprint_name(make("R1", fp="Lib:R_Plain")) == "R_Plain"


# --- sorting ------------------------------------------------------------------------


def test_smd_lines_come_first_then_through_hole_each_sorted_by_reference() -> None:
    parts = [
        make("J2", "Hdr", sym="Conn", fp=HEADER3, mpn="HDR-3"),
        make("C10", "1u", sym="Device:C", fp=C0603, mpn="CAP-B"),
        make("C2", "1u", sym="Device:C", fp=C0603, mpn="CAP-A"),
        make("RN1", "4x1k", sym="Device:R_Network", fp=C0603, mpn="ARRAY"),
        make("R12", "47k", mpn="OTHER-R"),
        make("J10", "Hdr", sym="Conn", fp=HEADER3, mpn="HDR-3"),
        make("A1", "Socket", sym="Conn", fp="Lib:Socket", mpn="SOCKET"),
    ]
    lines = lines_of(parts, through_hole={"J2", "J10", "A1"})
    assert [line.refs[0] for line in lines] == ["C2", "C10", "R12", "RN1", "A1", "J2"]
    assert [line.through_hole for line in lines] == [False] * 4 + [True] * 2
    assert [line.item for line in lines] == [1, 2, 3, 4, 5, 6]
    assert lines[-1].refs == ("J2", "J10")  # numbers sort as numbers: J2 before J10
    assert lines[-1].row()[-1] == "THT" and lines[0].row()[-1] == "SMD"


def test_total_parts_adds_every_quantity() -> None:
    parts = [make("R1"), make("R2"), make("C1", "1u", sym="Device:C", fp=C0603)]
    assert bom.total_parts(lines_of(parts)) == 3


# --- the default resistor table -----------------------------------------------------


def test_a_resistor_with_the_stand_in_part_number_gets_the_default_yageo_part() -> None:
    line = lines_of([make("R1", "4.7k")])[0]
    assert (line.mfr, line.mpn) == ("Yageo", "RC0603FR-074K7L")


def test_a_part_number_the_author_wrote_is_never_replaced_by_the_default() -> None:
    parts = [make("R1", "10k", mpn="ERJ-3EKF1002V", mfr="Panasonic"), make("R2", "10k")]
    mine, default = lines_of(parts)
    assert (mine.mfr, mine.mpn) == ("Panasonic", "ERJ-3EKF1002V")
    assert (default.mfr, default.mpn) == ("Yageo", "RC0603FR-0710KL")


def test_the_default_table_is_for_0603_resistors_only() -> None:
    wide = make("R1", "10k", fp="Resistor_SMD:R_0805_2012Metric")
    line = lines_of([wide])[0]
    assert line.mpn == "0603 10k 1%"  # left as it was, for the completeness check


def test_the_default_table_is_for_resistors_only() -> None:
    # a capacitor symbol in a resistor footprint, with no part number and a value the
    # table knows: only its symbol keeps it from getting a resistor's part number
    odd = make("C1", "10k", sym="Device:C", fp=R0603, mpn="")
    assert lines_of([odd])[0].mpn == ""
    assert lines_of([make("R1", "10k", mpn="")])[0].mpn == "RC0603FR-0710KL"


def test_the_stand_in_is_what_the_design_dsl_makes() -> None:
    design.reset()
    try:
        design.R("R1", "4.7k", "A", "B", "Block")
        made = design.PARTS[0]["mpn"]
    finally:
        design.reset()
    assert made == bom.stand_in_mpn("4.7k")


# --- overrides ----------------------------------------------------------------------


def test_an_mpn_table_entry_replaces_the_fields_it_names() -> None:
    parts = [make("J1", "H", sym="Conn", fp=HEADER3, mpn="HDR 1x3", desc="generic")]
    table = {"HDR 1x3": {"mfr": "Acme", "mpn": "AC-HDR-1X3", "desc": "Header 1x3"}}
    line = lines_of(parts, by_mpn=table)[0]
    assert (line.mfr, line.mpn, line.desc) == ("Acme", "AC-HDR-1X3", "Header 1x3")
    assert line.value == "H"  # a field the entry does not name is kept


def test_a_reference_entry_wins_over_an_mpn_entry_and_can_set_the_quantity() -> None:
    parts = [make("A1", "Socket", sym="Conn", fp="Lib:Socket", mpn="2x HDR")]
    line = lines_of(
        parts,
        by_mpn={"2x HDR": {"mpn": "FROM-MPN-TABLE"}},
        by_ref={"A1": {"mpn": "HDR-22", "mfr": "Acme", "qty": 2}},
    )[0]
    assert (line.mpn, line.mfr, line.qty) == ("HDR-22", "Acme", 2)


def test_a_quantity_override_counts_in_the_group_total() -> None:
    parts = [
        make("J1", "H", sym="Conn", fp=HEADER3, mpn="HDR"),
        make("J2", "H", sym="Conn", fp=HEADER3, mpn="HDR"),
    ]
    line = lines_of(parts, by_mpn={"HDR": {"qty": 2}})[0]
    assert line.qty == 4
    assert line.refs == ("J1", "J2")


def test_a_reference_entry_for_a_part_that_is_not_in_the_design_is_an_error() -> None:
    with pytest.raises(BomError, match=r"names 'A11'.*did you mean 'A1'"):
        lines_of([make("A1")], by_ref={"A11": {"qty": 2}})


def test_a_reference_entry_for_a_part_left_out_of_the_bom_is_not_an_error() -> None:
    parts = [make("R1"), make("R2", dnp=True)]
    assert len(lines_of(parts, by_ref={"R2": {"qty": 2}})) == 1


def test_an_mpn_entry_nobody_uses_is_ignored() -> None:
    assert len(lines_of([make("R1")], by_mpn={"UNUSED": {"mpn": "X"}})) == 1


# --- the line hook ------------------------------------------------------------------


def test_the_hook_sees_the_group_and_can_set_the_value_and_description() -> None:
    seen: list[Group] = []

    def hook(group: Group) -> dict[str, str]:
        seen.append(group)
        return {"value": f"{len(group.parts)} headers", "desc": "|".join(group.refs)}

    parts = [
        make("J2", "a", sym="Conn", fp=HEADER3, mpn="HDR"),
        make("J1", "b", sym="Conn", fp=HEADER3, mpn="HDR"),
    ]
    line = lines_of(parts, through_hole={"J1", "J2"}, line=hook)[0]
    assert (line.value, line.desc) == ("2 headers", "J2|J1")  # design order
    [group] = seen
    assert group.refs == ("J2", "J1")
    assert (group.value, group.mpn, group.qty, group.through_hole) == (
        "a",
        "HDR",
        2,
        True,
    )
    assert group.footprint == "PinHeader_1x03_P2.54mm_Vertical"
    assert [part["ref"] for part in group.parts] == ["J2", "J1"]


def test_a_hook_may_return_nothing_to_leave_a_line_alone() -> None:
    parts = [make("J1", "b", sym="Conn", fp=HEADER3, mpn="HDR")]
    line = lines_of(parts, line=lambda group: None)[0]
    assert (line.value, line.desc) == ("b", "Conn b")


@pytest.mark.parametrize(
    "returned",
    [{"mpn": "NOT-ALLOWED"}, {"value": 3}, ["value"], "value"],
)
def test_a_hook_that_returns_the_wrong_thing_is_an_error(returned: Any) -> None:
    parts = [make("J1", "b", sym="Conn", fp=HEADER3, mpn="HDR")]
    with pytest.raises(BomError, match="line"):
        lines_of(parts, line=lambda group: returned)


# --- completeness -------------------------------------------------------------------


def test_a_part_with_no_part_number_is_reported() -> None:
    parts = [make("C1", "100n", sym="Device:C", fp=C0603)]
    [problem] = bom.parts_without_mpn(parts)
    assert problem.ref == "C1"
    assert "no manufacturer part number" in problem.reason
    assert str(problem) == "C1 (100n): has no manufacturer part number"


def test_a_resistor_value_the_table_lacks_is_reported_as_the_stand_in() -> None:
    [problem] = bom.parts_without_mpn([make("R7", "4.3k")])
    assert problem.mpn == "0603 4.3k 1%"
    assert "stand-in" in problem.reason
    assert str(problem).startswith("R7 (4.3k): '0603 4.3k 1%' is the stand-in")


def test_a_part_number_with_a_space_in_it_is_reported_as_a_description() -> None:
    parts = [make("J1", "H", sym="Conn", fp=HEADER3, mpn="HDR 1x3 male")]
    [problem] = bom.parts_without_mpn(parts)
    assert "space" in problem.reason


def test_a_part_with_a_real_part_number_is_not_reported() -> None:
    parts = [
        make("U1", "IC", sym="Lib:IC", fp="Lib:SOT-23", mpn="SN74AHCT1G125DBVR"),
        make("D1", "LED", sym="Device:LED", fp="Lib:LED", mpn="150060RS75000"),
        make("R1", "10k"),  # the default table gives it one
    ]
    assert bom.parts_without_mpn(parts) == []


def test_completeness_judges_what_the_bom_will_show_after_the_overrides() -> None:
    parts = [
        make("J1", "H", sym="Conn", fp=HEADER3, mpn="HDR 1x3 male"),
        make("A1", "Socket", sym="Conn", fp="Lib:Socket", mpn=""),
    ]
    assert len(bom.parts_without_mpn(parts)) == 2
    fixed = Overrides(
        by_mpn={"HDR 1x3 male": {"mpn": "AC-HDR-1X3"}},
        by_ref={"A1": {"mpn": "AC-SOCKET-22"}},
    )
    assert bom.parts_without_mpn(parts, fixed) == []


def test_completeness_ignores_parts_that_are_not_fitted() -> None:
    parts = [
        make("C1", "100n", sym="Device:C", fp=C0603, dnp=True),
        make("C2", "100n", sym="Device:C", fp=C0603, in_bom=False),
        make("#PWR01", "x", sym="power:GND", fp="", mpn=""),
    ]
    assert bom.parts_without_mpn(parts) == []


def test_the_problems_are_sorted_by_reference_as_numbers() -> None:
    parts = [make(ref, "4.3k") for ref in ("R10", "R2", "C1")]
    parts[2] = make("C1", "1u", sym="Device:C", fp=C0603)
    assert [p.ref for p in bom.parts_without_mpn(parts)] == ["C1", "R2", "R10"]


def test_completeness_also_checks_the_reference_table() -> None:
    with pytest.raises(BomError, match="names 'Z9'"):
        bom.parts_without_mpn([make("R1")], Overrides(by_ref={"Z9": {"qty": 2}}))


# --- the files ----------------------------------------------------------------------


def sample_lines() -> list[bom.BomLine]:
    """Return a small BOM: a resistor pair, a capacitor with a comma, a header."""
    parts = [
        make("R1", "10k"),
        make("R2", "10k"),
        make(
            "C1",
            "1u, 25V",
            sym="Device:C",
            fp=C0603,
            mpn="CAP-1",
            mfr="Acme",
            desc='Cap "1u"',
        ),
        make("J1", "Hdr", sym="Conn", fp=HEADER3, mpn="HDR-3", mfr="Acme"),
    ]
    return lines_of(parts, through_hole={"J1"})


def test_the_csv_has_a_header_then_a_row_per_line_with_crlf_endings(
    tmp_path: Path,
) -> None:
    path = tmp_path / "bom.csv"
    bom.write_csv(path, sample_lines())
    raw = path.read_bytes()
    assert raw.startswith(
        b"Item,Qty,Designator,Manufacturer,Manufacturer Part Number,Value,"
        b"Description,Package / Footprint,Type\r\n"
    )
    assert raw.count(b"\r\n") == 4 and raw.count(b"\n") == 4  # no bare line feeds
    with open(path, newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    assert rows[1] == [
        "1", "1", "C1", "Acme", "CAP-1", "1u, 25V", 'Cap "1u"',
        "C_0603_1608Metric", "SMD",
    ]  # fmt: skip
    assert rows[2][:5] == ["2", "2", "R1,R2", "Yageo", "RC0603FR-0710KL"]
    assert rows[3][-1] == "THT"


def read_sheet(path: Path) -> list[list[Any]]:
    """Return every cell of the workbook's only sheet, row by row."""
    book = load_workbook(path)
    assert book.sheetnames == ["BOM"]
    return [list(row) for row in book["BOM"].iter_rows(values_only=True)]


def test_the_workbook_holds_the_same_cells_as_the_csv(tmp_path: Path) -> None:
    lines = sample_lines()
    path = tmp_path / "bom.xlsx"
    bom.write_xlsx(path, lines)
    cells = read_sheet(path)
    assert cells[0] == list(bom.HEADER)
    assert cells[1:] == [line.row() for line in lines]  # numbers stay numbers


def test_the_workbook_footer_follows_one_empty_row(tmp_path: Path) -> None:
    lines = sample_lines()
    path = tmp_path / "bom.xlsx"
    bom.write_xlsx(path, lines, ["J9 is a wire pad.", "H1 is a mounting hole."])
    cells = read_sheet(path)
    table = len(lines) + 1  # the header and the lines
    assert cells[table] == [None] * len(bom.HEADER)
    assert cells[table + 1][0] == "Not assembled / not in BOM:"
    assert cells[table + 2][0] == "J9 is a wire pad."
    assert cells[table + 3][0] == "H1 is a mounting hole."
    assert len(cells) == table + 4


def test_the_workbook_has_no_footer_when_the_project_gave_no_text(
    tmp_path: Path,
) -> None:
    lines = sample_lines()
    path = tmp_path / "bom.xlsx"
    bom.write_xlsx(path, lines)
    assert len(read_sheet(path)) == len(lines) + 1


# --- reading the board --------------------------------------------------------------


def board_text(*footprints: str) -> str:
    """Return the text of a board that holds the given footprint nodes."""
    return "(kicad_pcb (version 20260206)\n" + "\n".join(footprints) + "\n)\n"


def footprint(ref: str, *pad_types: str) -> str:
    """Return one footprint node with a pad of each type, in KiCad 10's layout."""
    pads = "".join(f'(pad "{n}" {t} circle (at 0 0))' for n, t in enumerate(pad_types))
    return f'(footprint "Lib:Name" (property "Reference" "{ref}") {pads})'


def test_a_footprint_with_a_plated_pad_is_through_hole(tmp_path: Path) -> None:
    path = tmp_path / "b.kicad_pcb"
    path.write_text(
        board_text(
            footprint("J1", "thru_hole", "thru_hole"),
            footprint("R1", "smd", "smd"),
            footprint("SW1", "smd", "np_thru_hole"),  # a locating peg is no THT
            footprint("MH1", "np_thru_hole"),
            footprint("J2", "smd", "thru_hole"),
        ),
        encoding="utf-8",
    )
    assert bom.board_footprints(path) == {
        "J1": True,
        "R1": False,
        "SW1": False,
        "MH1": False,
        "J2": True,
    }


def test_a_board_with_no_footprints_is_an_error(tmp_path: Path) -> None:
    path = tmp_path / "b.kicad_pcb"
    path.write_text(board_text(), encoding="utf-8")
    with pytest.raises(BomError, match="no footprints"):
        bom.board_footprints(path)


# --- the project's bom.py -----------------------------------------------------------


def project_with(tmp_path: Path, source: str | None) -> Path:
    """Return a folder holding a bom.py with ``source``, or none when it is None."""
    if source is not None:
        write_file(tmp_path / "bom.py", source)
    return tmp_path


def test_a_project_with_no_bom_py_has_no_overrides(tmp_path: Path) -> None:
    assert bom.load_overrides(project_with(tmp_path, None)) == Overrides()


def test_bom_py_names_are_read(tmp_path: Path) -> None:
    root = project_with(
        tmp_path,
        """
        MPN_OVERRIDE = {"HDR 1x3": {"mpn": "AC-HDR-1X3", "mfr": "Acme"}}
        REF_OVERRIDE = {"A1": {"qty": 2}}
        NOT_IN_BOM = ["J9 is a wire pad."]

        def line(group):
            return {"value": "Header"}
        """,
    )
    found = bom.load_overrides(root)
    assert found.by_mpn == {"HDR 1x3": {"mpn": "AC-HDR-1X3", "mfr": "Acme"}}
    assert found.by_ref == {"A1": {"qty": 2}}
    assert found.not_in_bom == ("J9 is a wire pad.",)
    assert found.line is not None
    assert found.line(None) == {"value": "Header"}  # type: ignore[arg-type]


@pytest.mark.parametrize(
    ("source", "message"),
    [
        ('MPN_OVERRIDE = ["x"]', "MPN_OVERRIDE in bom.py should be a dict"),
        ('REF_OVERRIDE = {"A1": 2}', "should map a reference"),
        ('REF_OVERRIDE = {1: {"qty": 2}}', "should map a reference"),
        ('MPN_OVERRIDE = {"x": {"mpm": "a"}}', "did you mean 'mpn'"),
        ('MPN_OVERRIDE = {"x": {"colour": "a"}}', "does not know"),
        ('MPN_OVERRIDE = {"x": {"qty": 0}}', "whole number >= 1"),
        ('MPN_OVERRIDE = {"x": {"qty": True}}', "whole number >= 1"),
        ('MPN_OVERRIDE = {"x": {"qty": 2.0}}', "whole number >= 1"),
        ('MPN_OVERRIDE = {"x": {"mpn": 5}}', "should be a string"),
        ("NOT_IN_BOM = 'one line'", "NOT_IN_BOM in bom.py should be a list"),
        ("NOT_IN_BOM = [1]", "NOT_IN_BOM in bom.py should be a list"),
        ("line = 3", "line in bom.py should be a function"),
        ("MPN_OVERIDE = {}", "did you mean MPN_OVERRIDE"),
        ("REF_OVERRIDES = {}", "did you mean REF_OVERRIDE"),
    ],
)
def test_a_mistake_in_bom_py_is_reported_with_what_to_fix(
    tmp_path: Path, source: str, message: str
) -> None:
    with pytest.raises(BomError, match=message):
        bom.load_overrides(project_with(tmp_path, source))


def test_other_upper_case_names_in_bom_py_are_the_users_own(tmp_path: Path) -> None:
    root = project_with(
        tmp_path,
        """
        HEADER = {"mpn": "AC-HDR-1X3"}
        MPN_OVERRIDE = {"HDR 1x3": HEADER}
        """,
    )
    assert bom.load_overrides(root).by_mpn == {"HDR 1x3": {"mpn": "AC-HDR-1X3"}}
