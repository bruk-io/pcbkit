"""The bill of materials: group the fitted parts into lines, write them, check them.

``bom_lines`` turns a design's parts into BOM lines: parts with the same manufacturer
part number and footprint share a line. It is a pure function of the parts, the set of
references that have through-hole pads, and the project's optional overrides, so it
needs no KiCad. ``write_csv`` and ``write_xlsx`` write the lines; ``parts_without_mpn``
says which fitted parts a buyer could not order.

A line takes its manufacturer, part number, description and value from the first part
of its group in the order design.py declared them, never from the first reference after
sorting. Lines are sorted surface-mount first, then by the first reference's letters
and number.

What a part is bought as can differ from what the schematic says (a generic header
named in design.py, bought as one real part). The project's optional ``bom.py`` holds
that: see ``load_overrides`` and docs/project-interface.md.
"""

from __future__ import annotations

import csv
import difflib
import re
from collections.abc import Callable, Collection, Mapping, Sequence
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any

import click

from pcbkit.design import Part
from pcbkit.kicad.sexp import findall, parse
from pcbkit.project import import_optional_project_module

HEADER = (
    "Item",
    "Qty",
    "Designator",
    "Manufacturer",
    "Manufacturer Part Number",
    "Value",
    "Description",
    "Package / Footprint",
    "Type",
)

# Yageo RC0603FR-07 series, 1 % 0603: the part number a resistor gets when design.py
# gave it none, by its value. A value that is not here is left for the user to name.
RESISTOR_MPN = {
    "10": "RC0603FR-0710RL",
    "220": "RC0603FR-07220RL",
    "330": "RC0603FR-07330RL",
    "100": "RC0603FR-07100RL",
    "1k": "RC0603FR-071KL",
    "3.3k": "RC0603FR-073K3L",
    "1.5k": "RC0603FR-071K5L",
    "2.2k": "RC0603FR-072K2L",
    "4.7k": "RC0603FR-074K7L",
    "10k": "RC0603FR-0710KL",
    "18k": "RC0603FR-0718KL",
    "22k": "RC0603FR-0722KL",
    "100k": "RC0603FR-07100KL",
    "300": "RC0603FR-07300RL",
    "560": "RC0603FR-07560RL",
    "5.23k": "RC0603FR-075K23L",
    "30.1k": "RC0603FR-0730K1L",
}
RESISTOR_MFR = "Yageo"

# The widths of the workbook's nine columns, in characters.
_WIDTHS = (6, 5, 40, 20, 26, 12, 60, 42, 6)

# What a project's bom.py may define, and the fields each entry or hook may set.
OVERRIDE_NAMES = ("MPN_OVERRIDE", "REF_OVERRIDE", "NOT_IN_BOM", "line")
PART_FIELDS = ("mfr", "mpn", "desc", "qty")
LINE_FIELDS = ("value", "desc")

_REFERENCE = re.compile(r"([A-Z]+)(\d+)")


class BomError(click.ClickException):
    """Report a mistake in a project's bom.py or a part the BOM cannot place."""


@dataclass(frozen=True)
class Group:
    """The parts that make one BOM line, as a project's ``line`` hook sees them.

    ``parts`` are in design.py order. ``mfr``, ``mpn``, ``value`` and ``desc`` are the
    first part's, after the overrides, and ``qty`` is the whole group's.
    """

    parts: tuple[Part, ...]
    footprint: str
    mfr: str
    mpn: str
    value: str
    desc: str
    qty: int
    through_hole: bool

    @property
    def refs(self) -> tuple[str, ...]:
        """Return the references of the parts, in design.py order."""
        return tuple(part["ref"] for part in self.parts)


LineHook = Callable[[Group], "Mapping[str, str] | None"]


@dataclass(frozen=True)
class Overrides:
    """What a project's bom.py says about how parts are bought and shown.

    ``by_mpn`` maps an MPN as written in design.py to the fields to replace;
    ``by_ref`` does the same for one reference, and wins. ``line`` may replace the
    value and description a finished line shows. ``not_in_bom`` is the text the
    workbook carries under its table.
    """

    by_mpn: Mapping[str, Mapping[str, Any]] = field(default_factory=dict)
    by_ref: Mapping[str, Mapping[str, Any]] = field(default_factory=dict)
    line: LineHook | None = None
    not_in_bom: tuple[str, ...] = ()


@dataclass(frozen=True)
class BomLine:
    """One line of the bill of materials."""

    item: int
    qty: int
    refs: tuple[str, ...]
    mfr: str
    mpn: str
    value: str
    desc: str
    footprint: str
    through_hole: bool

    def row(self) -> list[Any]:
        """Return the line as the nine cells of a BOM row, in ``HEADER`` order."""
        return [
            self.item,
            self.qty,
            ",".join(self.refs),
            self.mfr,
            self.mpn,
            self.value,
            self.desc,
            self.footprint,
            "THT" if self.through_hole else "SMD",
        ]


@dataclass(frozen=True)
class MpnProblem:
    """A fitted part that cannot be ordered, and why."""

    ref: str
    value: str
    mpn: str
    reason: str

    def __str__(self) -> str:
        """Return "R4 (10k): reason", with the part number shown when there is one."""
        shown = f" {self.mpn!r}" if self.mpn else ""
        return f"{self.ref} ({self.value}):{shown} {self.reason}"


@dataclass(frozen=True)
class _Resolved:
    """What one part is bought as, once the overrides and defaults have applied."""

    mfr: str
    mpn: str
    desc: str
    qty: int


# --- the parts ---------------------------------------------------------------------


def refkey(ref: str) -> tuple[str, int]:
    """Return a sort key for a reference: its letters, then its number."""
    found = _REFERENCE.match(ref)
    return (found.group(1), int(found.group(2))) if found else (ref, 0)


def fitted(parts: Sequence[Part]) -> list[Part]:
    """Return the parts that belong in the BOM: not flags, not DNP, not bom=False."""
    return [
        part
        for part in parts
        if part["bom"] and not part["ref"].startswith("#") and not part["dnp"]
    ]


def footprint_name(part: Part) -> str:
    """Return the part's footprint without its library: "R_0603_1608Metric"."""
    return part["fp"].split(":", 1)[-1]


def stand_in_mpn(value: str) -> str:
    """Return the part number ``pcbkit.design.R`` makes up for ``value``."""
    return f"0603 {value} 1%"


def _resolve(part: Part, overrides: Overrides) -> _Resolved:
    """Return what ``part`` is bought as.

    The steps, in order: the project's MPN table; the default resistor table, only for
    a 0603 ``Device:R`` still carrying the helper's stand-in part number (a part number
    the author wrote is never replaced); the project's table for this reference.
    """
    mfr, mpn, desc, qty = part["mfr"], part["mpn"], part["desc"], 1
    found = overrides.by_mpn.get(part["mpn"])
    if found is not None:
        mfr = found.get("mfr", mfr)
        mpn = found.get("mpn", mpn)
        desc = found.get("desc", desc)
        qty = found.get("qty", qty)
    if (
        part["sym"] == "Device:R"
        and mpn in ("", stand_in_mpn(part["value"]))
        and footprint_name(part).startswith("R_0603")
        and part["value"] in RESISTOR_MPN
    ):
        mfr, mpn = RESISTOR_MFR, RESISTOR_MPN[part["value"]]
    found = overrides.by_ref.get(part["ref"])
    if found is not None:
        mfr = found.get("mfr", mfr)
        mpn = found.get("mpn", mpn)
        desc = found.get("desc", desc)
        qty = found.get("qty", qty)
    return _Resolved(mfr, mpn, desc, qty)


def _check_refs(parts: Sequence[Part], overrides: Overrides) -> None:
    """Raise BomError for a ``REF_OVERRIDE`` key that is not in the design."""
    known = {part["ref"] for part in parts}
    for ref in overrides.by_ref:
        if ref not in known:
            close = difflib.get_close_matches(ref, sorted(known), n=1)
            hint = f" (did you mean {close[0]!r}?)" if close else ""
            raise BomError(
                f"REF_OVERRIDE in bom.py names {ref!r}, which is not in the "
                f"design{hint}"
            )


# --- the lines ---------------------------------------------------------------------


def bom_lines(
    parts: Sequence[Part],
    through_hole: Collection[str],
    overrides: Overrides | None = None,
) -> list[BomLine]:
    """Group the fitted ``parts`` into BOM lines, sorted and numbered.

    ``through_hole`` holds the references whose footprints have plated through-hole
    pads (``board_footprints`` reads them from the saved board): their lines are Type
    THT, come after every SMD line, and the rest are SMD. Parts that share a part
    number and footprint share a line; a part with no part number is grouped by its
    value instead. Raise BomError if ``overrides`` names a reference that is not in the
    design, or a hook returns something other than a mapping of ``LINE_FIELDS``.
    """
    over = overrides or Overrides()
    _check_refs(parts, over)
    groups: dict[tuple[str, str], list[tuple[Part, _Resolved]]] = {}
    for part in fitted(parts):
        bought = _resolve(part, over)
        footprint = footprint_name(part)
        key = (bought.mpn, footprint) if bought.mpn else (part["value"], footprint)
        groups.setdefault(key, []).append((part, bought))
    lines: list[BomLine] = []
    for (_, footprint), members in groups.items():
        first_part, first = members[0]
        group = Group(
            parts=tuple(part for part, _ in members),
            footprint=footprint,
            mfr=first.mfr,
            mpn=first.mpn,
            value=first_part["value"],
            desc=first.desc,
            qty=sum(bought.qty for _, bought in members),
            through_hole=first_part["ref"] in through_hole,
        )
        shown = _shown(group, over.line)
        lines.append(
            BomLine(
                item=0,
                qty=group.qty,
                refs=tuple(sorted(group.refs, key=refkey)),
                mfr=group.mfr,
                mpn=group.mpn,
                value=shown.get("value", group.value),
                desc=shown.get("desc", group.desc),
                footprint=footprint,
                through_hole=group.through_hole,
            )
        )
    lines.sort(key=lambda line: (line.through_hole, refkey(line.refs[0])))
    return [replace(line, item=number) for number, line in enumerate(lines, start=1)]


def _shown(group: Group, hook: LineHook | None) -> Mapping[str, str]:
    """Return the value and description the project's hook sets for ``group``."""
    if hook is None:
        return {}
    said = hook(group)
    if said is None:
        return {}
    if not isinstance(said, Mapping):
        raise BomError(
            f"line() in bom.py should return None or a dict, got {said!r} "
            f"for {', '.join(group.refs)}"
        )
    for name, text in said.items():
        if name not in LINE_FIELDS or not isinstance(text, str):
            raise BomError(
                f"line() in bom.py returned {name!r}: {text!r} for "
                f"{', '.join(group.refs)}; it may set {' and '.join(LINE_FIELDS)}, "
                "each a string"
            )
    return said


def total_parts(lines: Sequence[BomLine]) -> int:
    """Return the number of physical parts: the sum of every line's quantity."""
    return sum(line.qty for line in lines)


# --- completeness ------------------------------------------------------------------


def parts_without_mpn(
    parts: Sequence[Part], overrides: Overrides | None = None
) -> list[MpnProblem]:
    """Return the fitted parts a buyer could not order, sorted by reference.

    This runs on what each part is bought as, after the overrides and the default
    resistor table, because that is what lands in the BOM. A part is listed when its
    part number is empty, is still the stand-in that ``R`` makes up for a resistor,
    or has a space in it. The last is a heuristic: a real manufacturer part number is
    one word, and one with a space is a description such as "pin header 1x3 male"; a
    real part number that does contain a space can only be reported, not accepted.
    """
    over = overrides or Overrides()
    _check_refs(parts, over)
    found: list[MpnProblem] = []
    for part in fitted(parts):
        mpn = _resolve(part, over).mpn
        reason = mpn_problem(mpn, part["value"])
        if reason is not None:
            found.append(MpnProblem(part["ref"], part["value"], mpn, reason))
    return sorted(found, key=lambda problem: refkey(problem.ref))


def mpn_problem(mpn: str, value: str) -> str | None:
    """Return why a buyer could not order ``mpn`` (a part of ``value``), or None.

    The rules of ``parts_without_mpn``, for one part number: empty, the stand-in that
    ``R`` makes up for a resistor of that value, or a description with a space in it.
    """
    if not mpn.strip():
        return "has no manufacturer part number"
    if mpn == stand_in_mpn(value):
        return (
            "is the stand-in part number pcbkit.design.R makes up: give the "
            "part an mpn, or its value a default resistor"
        )
    if re.search(r"\s", mpn):
        return "has a space in it, so it is a description, not a part number"
    return None


# --- the board ---------------------------------------------------------------------


def board_footprints(pcb: Path) -> dict[str, bool]:
    """Return every footprint on a saved board: its reference, and whether it has a
    plated through-hole pad.

    A footprint is through-hole when one of its pads is of type ``thru_hole``; a part
    with only SMD pads, or only unplated holes (locating pegs), is surface-mount.
    Raises BomError if the file holds no footprint with a reference (an empty or
    unreadable board).
    """
    tree = parse(Path(pcb).read_text(encoding="utf-8"))
    found: dict[str, bool] = {}
    for footprint in findall(tree, "footprint"):
        ref = next(
            (
                str(prop[2])
                for prop in findall(footprint, "property")
                if len(prop) > 2 and prop[1] == "Reference"
            ),
            "",
        )
        if ref:
            found[ref] = any(
                len(pad) > 2 and pad[2] == "thru_hole"
                for pad in findall(footprint, "pad")
            )
    if not found:
        raise BomError(f"{pcb} has no footprints: run `pcbkit build` first")
    return found


# --- writing -----------------------------------------------------------------------


def write_csv(path: Path, lines: Sequence[BomLine]) -> None:
    """Write the BOM as CSV: a header row, then one row per line, with CRLF endings."""
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADER)
        writer.writerows(line.row() for line in lines)


def read_csv(path: Path) -> list[BomLine]:
    """Read a BOM CSV that ``write_csv`` wrote back into lines.

    Raise BomError if the file is not one: a wrong header, or a row that is short or
    has a quantity or type it should not.
    """
    with open(path, newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    if not rows or tuple(rows[0]) != HEADER:
        raise BomError(f"{path} is not a pcbkit BOM: its header is not {list(HEADER)}")
    lines = []
    for number, row in enumerate(rows[1:], start=2):
        if len(row) != len(HEADER) or row[8] not in ("THT", "SMD"):
            raise BomError(f"{path} row {number} is not a BOM line: {row}")
        try:
            item, qty = int(row[0]), int(row[1])
        except ValueError:
            raise BomError(f"{path} row {number} has no whole Item and Qty") from None
        lines.append(
            BomLine(
                item=item,
                qty=qty,
                refs=tuple(row[2].split(",")),
                mfr=row[3],
                mpn=row[4],
                value=row[5],
                desc=row[6],
                footprint=row[7],
                through_hole=row[8] == "THT",
            )
        )
    return lines


def write_xlsx(
    path: Path, lines: Sequence[BomLine], not_in_bom: Sequence[str] = ()
) -> None:
    """Write the BOM as a workbook with one sheet, "BOM".

    Under the table, after one empty row, goes a bold "Not assembled / not in BOM:"
    heading and then each text of ``not_in_bom`` on a row of its own; with no texts
    there is neither the heading nor the rows.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "BOM"
    sheet.append(list(HEADER))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F3B4C")
    for line in lines:
        sheet.append(line.row())
    for index, width in enumerate(_WIDTHS):
        sheet.column_dimensions[chr(65 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    if not_in_bom:
        first = sheet.max_row + 2
        sheet.cell(first, 1, "Not assembled / not in BOM:").font = Font(bold=True)
        for offset, text in enumerate(not_in_bom, start=1):
            sheet.cell(first + offset, 1, text)
    book.save(path)


# --- the project's bom.py ----------------------------------------------------------


def _table(module: Any, name: str) -> dict[str, dict[str, Any]]:
    """Return the module's override table called ``name``; check every entry."""
    table = getattr(module, name, {})
    if not isinstance(table, dict):
        raise BomError(f"{name} in bom.py should be a dict of {_KEYS[name]} to fields")
    for key, entry in table.items():
        if not isinstance(key, str) or not isinstance(entry, dict):
            raise BomError(
                f"{name} in bom.py should map {_KEYS[name]} (a string) to a dict of "
                f"{', '.join(PART_FIELDS)}; got {key!r}: {entry!r}"
            )
        for fieldname, value in entry.items():
            if fieldname not in PART_FIELDS:
                close = difflib.get_close_matches(str(fieldname), PART_FIELDS, n=1)
                hint = f" (did you mean {close[0]!r}?)" if close else ""
                raise BomError(
                    f"{name}[{key!r}] in bom.py has the field {fieldname!r}, which "
                    f"pcbkit does not know{hint}; the fields are "
                    f"{', '.join(PART_FIELDS)}"
                )
            if fieldname == "qty":
                ok = (
                    isinstance(value, int) and not isinstance(value, bool) and value > 0
                )
                want = "a whole number >= 1"
            else:
                ok = isinstance(value, str)
                want = "a string"
            if not ok:
                raise BomError(
                    f"{name}[{key!r}][{fieldname!r}] in bom.py should be {want}, "
                    f"got {value!r}"
                )
    return table


_KEYS = {"MPN_OVERRIDE": "an MPN written in design.py", "REF_OVERRIDE": "a reference"}


def load_overrides(root: Path) -> Overrides:
    """Return what the ``bom.py`` in the project at ``root`` defines; nothing if absent.

    Raise BomError for a name of the wrong type, an unknown field, or an upper-case
    name that looks like a misspelling of one pcbkit reads (``MPN_OVERIDE``).
    """
    module = import_optional_project_module(Path(root), "bom")
    if module is None:
        return Overrides()
    for name in vars(module):
        if name.isupper() and name not in OVERRIDE_NAMES:
            close = difflib.get_close_matches(name, OVERRIDE_NAMES, n=1, cutoff=0.75)
            if close:
                raise BomError(
                    f"bom.py defines {name}, which pcbkit does not read: "
                    f"did you mean {close[0]}?"
                )
    notes = getattr(module, "NOT_IN_BOM", [])
    if not isinstance(notes, (list, tuple)) or not all(
        isinstance(text, str) for text in notes
    ):
        raise BomError("NOT_IN_BOM in bom.py should be a list of strings")
    hook = getattr(module, "line", None)
    if hook is not None and not callable(hook):
        raise BomError("line in bom.py should be a function taking one Group")
    return Overrides(
        by_mpn=_table(module, "MPN_OVERRIDE"),
        by_ref=_table(module, "REF_OVERRIDE"),
        line=hook,
        not_in_bom=tuple(notes),
    )
